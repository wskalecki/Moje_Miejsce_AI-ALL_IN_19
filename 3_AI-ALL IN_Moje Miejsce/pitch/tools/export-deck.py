"""Build deck/export/moje-miejsce-pitch.pdf and .pptx from the slide screenshots (run deck-screens.mjs first).

The PPTX keeps the speaker notes and embeds the system presentation video on the video slide.
Requires: pip install python-pptx pillow
"""
import pathlib, re
from html import unescape
from PIL import Image
from pptx import Presentation
from pptx.util import Emu

root = pathlib.Path(__file__).resolve().parent.parent / "deck"
exp = root / "export"
slides = sorted(exp.glob("slide-*.png"), key=lambda p: int(re.findall(r"\d+", p.stem)[0]))
html = (root / "index.html").read_text()
sections = re.findall(r'<section class="slide[^"]*"(.*?)</section>', html, re.S)
notes = [unescape(re.sub(r"\s+", " ", re.sub(r"<[^>]+>", "", m.group(1)))).strip() if (m := re.search(r'<div class="notes">(.*?)</div>', s, re.S)) else "" for s in sections]
video_idx = next(i for i, s in enumerate(sections) if "data-video" in s.split(">")[0])

# PDF
imgs = [Image.open(p).convert("RGB") for p in slides]
imgs[0].save(exp / "moje-miejsce-pitch.pdf", save_all=True, append_images=imgs[1:], resolution=144)

# PPTX
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
blank = prs.slide_layouts[6]
for i, p in enumerate(slides):
    s = prs.slides.add_slide(blank)
    if i == video_idx:
        s.shapes.add_movie(str(root / "media" / "moje-miejsce-system.mp4"), 0, 0, prs.slide_width, prs.slide_height,
                           poster_frame_image=str(root / "img" / "video-poster.jpg"), mime_type="video/mp4")
    else:
        s.shapes.add_picture(str(p), 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = notes[i] if i < len(notes) else ""
prs.save(exp / "moje-miejsce-pitch.pptx")
print("pdf + pptx →", exp)
